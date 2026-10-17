"""Synthetic DOCX/PDF extraction pilot; no paid services or corpus publication."""

import argparse
import json
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from chunking import chunk_document
from documents import read_docx, read_pdf_with_pdfplumber, read_pdf_with_pypdf
from extraction_quality import document_diagnostics


def make_pdf(path):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    lines = [
        "1 w",
        "50 600 m 450 600 l S",
        "50 570 m 450 570 l S",
        "50 540 m 450 540 l S",
        "50 510 m 450 510 l S",
        "50 510 m 50 600 l S",
        "250 510 m 250 600 l S",
        "450 510 m 450 600 l S",
    ]
    for x, y, text in [
        (60, 580, "Species"),
        (260, 580, "Fee (INR)"),
        (60, 550, "Timber"),
        (260, 550, "1,000"),
        (60, 520, "Bamboo"),
        (260, 520, "0 (exempt)"),
        (50, 480, "Bamboo still requires approval before transport."),
    ]:
        lines.append(f'BT /F1 12 Tf {x} {y} Td ({text.replace("(", "[").replace(")", "]")}) Tj ET')
    stream = DecodedStreamObject()
    stream.set_data("\n".join(lines).encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as output:
        writer.write(output)


def run_pilot():
    with TemporaryDirectory(prefix="answer-extraction-") as directory:
        folder = Path(directory)
        document = Document()
        document.add_heading("1. Transport requirements", level=1)
        document.add_heading("1.1 Fee schedule", level=2)
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Species"
        table.rows[0].cells[1].text = "Fee (INR)"
        for species, fee in [("Timber", "1,000"), ("Bamboo", "0 (exempt)")]:
            cells = table.add_row().cells
            cells[0].text, cells[1].text = species, fee
        document.add_paragraph("Bamboo still requires approval before transport.")
        document.add_paragraph("वन परिवहन अनुमति आवश्यक है।")
        docx = folder / "synthetic.docx"
        document.save(docx)
        pdf = folder / "synthetic.pdf"
        make_pdf(pdf)
        results = []
        for method, pages in [
            ("docx", read_docx(docx)),
            ("pdf-native", read_pdf_with_pypdf(pdf)),
            ("pdf-table-aware", read_pdf_with_pdfplumber(pdf)),
        ]:
            doc = {
                "source": "synthetic." + ("docx" if method == "docx" else "pdf"),
                "title": "Synthetic schedule",
                "kind": "document",
                "pages": pages,
                "metadata": {},
            }
            chunks = chunk_document(doc)
            text = "\n".join(c["content"] for c in chunks)
            required = ["Timber", "1,000", "Bamboo", "approval before transport"]
            if method == "docx":
                required.append("वन परिवहन अनुमति आवश्यक है।")
            results.append(
                {
                    "method": method,
                    "checks": {term: term in text for term in required},
                    "chunk_count": len(chunks),
                    "table_units": sum("table" in c["metadata"]["unit_types"] for c in chunks),
                    "diagnostics": document_diagnostics(pages),
                }
            )
    return {"scope": "Synthetic known-content extraction checks; not a production-document review", "results": results}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    report = run_pilot()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n")
    return 0 if all(all(r["checks"].values()) for r in report["results"]) else 1


if __name__ == "__main__":
    raise SystemExit(main())
