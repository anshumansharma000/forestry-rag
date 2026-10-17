from pathlib import Path
from datetime import date
import calendar, json
from decimal import Decimal, ROUND_HALF_UP
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = Path(__file__).resolve().parent
D = Decimal
def money(v):
    return f'{D(str(v)).quantize(D("1"), rounding=ROUND_HALF_UP):,}'
def lakh(v): return f'INR {float(v)/100000:.2f} lakh'
FIXED = D('15725')
rows=[]
for i in range(12):
    year=2026+(10+i)//12
    month=(10+i)%12+1
    days=calendar.monthrange(year,month)[1]
    queries=200*D('4.5')*days
    rate=D('.77') if year==2026 else D('1.40525')
    ai=queries*rate
    subtotal=ai+FIXED
    tax=subtotal*D('.18')
    cash=subtotal+tax
    rows.append(dict(month=date(year,month,1).strftime('%b %Y'),days=days,queries=queries,ai=ai,other=FIXED,tax=tax,cash=cash,reserve=cash*D('.15'),protected=cash*D('1.15'),allocation=D('50000') if year==2026 else D('75000')))
tot={k:sum(r[k] for r in rows) for k in ['queries','ai','other','tax','cash','reserve','protected','allocation']}
def scenario(dau, queries_per_day=D('4.5'), mult=D(1)):
    ai=sum(D(dau)*queries_per_day*r['days']*(D('.77') if i<2 else D('1.40525'))*mult for i,r in enumerate(rows))
    cash=(ai+FIXED*12)*D('1.18')
    return cash,cash*D('1.15')+D(10000)

doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.27); s.page_height=Inches(11.69)
s.top_margin=Inches(.67);s.bottom_margin=Inches(.65)
s.left_margin=Inches(.7);s.right_margin=Inches(.7)
s.header_distance=Inches(.25);s.footer_distance=Inches(.27)
for name in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(7)
    for border in list(st.element.xpath('./w:pPr/w:pBdr')):
        border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(27)
doc.styles['Title'].paragraph_format.space_after=Pt(10)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Aranyabodh  |  Operating cost projection  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Aranyabodh Annual Operating Cost Projection'
doc.core_properties.subject='November 2026 to October 2027 operating cost plan for 1000 registered users'
doc.core_properties.author='Aranyabodh'

def p(text='',bold=False,size=None):
    para=doc.add_paragraph();run=para.add_run(text);run.bold=bold
    if size:run.font.size=Pt(size)
    return para
def h(text):doc.add_heading(text,1)
def h2(text):doc.add_heading(text,2)
def newpage():doc.add_page_break()
def table(headers,data,widths,size=9.3):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for j,w in enumerate(widths):t.columns[j].width=Inches(w)
    for j,txt in enumerate(headers):t.rows[0].cells[j].text=txt
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in data:
        cells=t.add_row().cells
        for j,v in enumerate(values):cells[j].text=str(v)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'263F50' if i==0 else ('F1F4F6' if i%2 else 'FFFFFF'));tcPr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcPr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge,val in [('top','75'),('bottom','75'),('left','95'),('right','95')]:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),val);el.set(qn('w:type'),'dxa');margins.append(el)
            tcPr.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.space_before=Pt(0);para.paragraph_format.line_spacing=1.04
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.RIGHT
                if len(headers)==2 and j==1:para.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    run.font.size=Pt(size);run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000');run.bold=i==0
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def link(label,url):
    para=doc.add_paragraph();para.paragraph_format.space_after=Pt(7)
    el=OxmlElement('w:hyperlink');el.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');prop=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'245A75');prop.append(color)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');prop.append(sz);run.append(prop)
    txt=OxmlElement('w:t');txt.text=label;run.append(txt);el.append(run);para._p.append(el)

# Page 1
doc.add_paragraph('Aranyabodh Annual Operating Cost Projection','Title')
p('November 2026 to October 2027',bold=True,size=14)
p('Prepared 15 October 2026  |  Currency INR  |  Planning population 1,000 registered users',size=10)
h('Executive assessment')
p(f'Aranyabodh should reserve INR 8.60 lakh for the next 12 full operating months. The base projection totals {lakh(tot["cash"])} in recurring cash requirements, including a provisional tax allowance. Adding 15% contingency and a separate INR 10,000 initial onboarding allowance produces a calculated funding requirement of {lakh(tot["protected"]+10000)}. The rounded funding recommendation provides a further margin.')
table(['Annual funding measure','INR'],[
    ['Recurring services and operating allowances before tax',money(tot['ai']+tot['other'])],
    ['Provisional tax allowance at 18%',money(tot['tax'])],
    ['Recurring cash requirement including tax allowance',money(tot['cash'])],
    ['Contingency at 15% of recurring cash requirement',money(tot['reserve'])],
    ['Initial onboarding envelope separate from recurring costs','10,000'],
    ['Calculated first year funding requirement',money(tot['protected']+10000)],
    ['Recommended first year funding allocation','860,000'],
],[4.9,1.95],10)
p('Funding schedule: INR 50,000 per month in November and December 2026, followed by INR 75,000 per month from January to October 2027. Add INR 10,000 in November for initial provisioning and document ingestion. The annual allocation is therefore INR 860,000.',bold=True)
h2('Purpose and cost boundary')
p('This report supports operating-budget approval for a forest-document question answering service used by forest department officers and commercial entities. It covers the app backend, document processing, database and retrieval, model APIs, storage, backups, communications, monitoring and modest business-tool subscriptions.')
p('Development, maintenance labour, salaries, manual document curation and sales labour are excluded. Subscription revenue is not offset against costs. Revenue-dependent payment processing and procurement-specific requirements are shown separately because no collection volume or procurement terms have been established.')
p('All figures are estimates. Published provider rates, project evidence and planning assumptions are identified separately. Contingency and optional-service provisions are funding allowances rather than confirmed invoice charges.',size=10)

# Page 2
newpage();h('1 Planning assumptions')
p('The model holds the registered population and average activity constant for the full year. This gives a steady operating baseline rather than an acquisition or adoption forecast. Calendar-day calculations capture the variation between months.')
table(['Assumption','Base case and evidence status'],[
    ['Registered population','1,000 users. User-specified total, not 1,000 daily users.'],
    ['Monthly active users','600 users or 60% of registered users. Planning assumption.'],
    ['Daily active users','200 users or 20% of registered users. Planning assumption.'],
    ['DAU to MAU ratio','33.3%. Equivalent to about 10 active days per MAU in a 30-day month.'],
    ['Questions per active day','4.5 per DAU. Midpoint of the user-specified range of 4 to 5.'],
    ['Monthly question volume','DAU × 4.5 × actual calendar days. 27,000 in a 30-day month.'],
    ['Annual question volume',f'{money(tot["queries"])} questions across 365 days.'],
    ['Question mix','25% simple, 50% procedural and 25% complex. Assumed.'],
    ['Initial corpus','500 documents × 20 pages = 10,000 pages. Assumed.'],
    ['Monthly additions','100 documents × 20 pages = 2,000 pages. Assumed.'],
    ['OCR share','50% of pages require OCR. Native-text files avoid OCR. Assumed.'],
    ['Planning exchange rate','INR 90 per USD. Budget convention, not a verified spot rate.'],
],[2.0,4.85],9.5)
h2('Why these assumptions are usable')
p('A specialist reference app is likely to be used when an officer or business needs to interpret a document or complete a procedure. The base case therefore assumes intermittent use by much of the population and repeated use by a smaller working group. It is not supported by a measured industry benchmark; the low and high scenarios test that uncertainty explicitly.')
p('The repo describes a Gemini-backed retrieval service using Supabase Postgres with pgvector, document ingestion and citation-based answers. This architecture supports the selected cost categories. Current source files do not establish the live production capacity, enabled integrations or actual user behaviour. [10]')
h2('Tax and exchange rate treatment')
p('An 18% tax cash allowance is applied to the complete operating subtotal for conservative budgeting. This is not an invoice-specific GST determination. Recoverable input credits, reverse-charge treatment and supplier differences are not netted off. The 15% contingency covers moderate currency movement and spending uncertainty; no separate foreign-exchange surcharge is added.')

# Page 3
newpage();h('2 AI usage and unit economics')
p('One user question can trigger several model calls for rewriting, planning, answering and verification. The following per-question allowances aggregate those generation stages. They do not represent the price of one API call. Thinking tokens are included in billable output. [1, 10]')
table(['USD per million tokens','Input','Output'],[
    ['Gemini 3.5 Flash Lite throughout projection','0.30','2.50'],
    ['Gemini 3.8 Flash through December 2026','0.75','3.75'],
    ['Gemini 3.8 Flash from January 2027','1.50','7.50'],
],[4.65,1.1,1.1],10)
p('Google publishes the Flash increase effective 1 January 2027. The projection retains that announced change and holds all other model rates constant. It assumes standard paid API usage without caching discounts, batch discounts or promotional credits. [1]')
table(['Question class','Share','2026 INR','Flash bill share','2027 INR'],[
    ['Simple','25%','0.100','0%','0.100'],
    ['Procedural','50%','0.600','80%','1.080'],
    ['Complex','25%','1.500','90%','2.850'],
],[1.65,.65,1.25,1.9,1.4],9.5)
p('Class costs and Flash bill shares are planning assumptions. The repo records selected procedural generation costs of approximately INR 0.296, 0.764 and 0.682, while other runs show substantially higher costs when reasoning or retries increase. These examples support a broad cost range but do not validate the assumed mix or a fixed maximum. [10]')
h2('Reproducible calculation')
p('2026 weighted generation cost = (25% × 0.10) + (50% × 0.60) + (25% × 1.50) = INR 0.70 per question.')
p('Add 10% for recovery attempts: INR 0.70 × 1.10 = INR 0.77000 per question.')
p('2027 class cost = 2026 class cost × (1 + assumed Flash share of its bill). The resulting weighted cost is INR 1.27750; including recovery attempts gives INR 1.40525 per question.')
p('For a 30-day base month: 200 × 4.5 × 30 = 27,000 questions. Generation is INR 20,790 through December and INR 37,941.75 from January. The 10% recovery allowance is already included and is not applied again elsewhere.')
h2('Other API consumption')
p('The INR 500 monthly OCR and embedding envelope assumes 1,000 scanned pages, 1.44 million document tokens after chunk overlap and 200 query-embedding tokens per question. At 27,000 questions, indicative costs are INR 135 OCR, INR 26 document embeddings and INR 97 query embeddings, leaving headroom. OCR free allowances are conservatively ignored; advanced extraction is not included. [1, 4]')

# Page 4
newpage();h('3 Detailed recurring cost breakdown')
p('All amounts below are INR before the provisional tax allowance. The two monthly columns use 30 days for comparison. The annual column uses 365 calendar days. Allowances are retained in the funded base case even if actual invoices are lower.')
costs=[
('AI generation','20,790','37,942',money(tot['ai'])),
('Backend API hosting','2,250','2,250','27,000'),
('Document processing worker','2,250','2,250','27,000'),
('Supabase Pro and Small compute','2,700','2,700','32,400'),
('Redis or managed job queue','900','900','10,800'),
('Storage backup and bandwidth envelope','1,000','1,000','12,000'),
('Monitoring logs and error tracking','1,000','1,000','12,000'),
('OCR and embedding envelope','500','500','6,000'),
('Email SMS and WhatsApp envelope','2,000','2,000','24,000'),
('Domain annual fee allocation','125','125','1,500'),
('Business and support subscriptions','2,000','2,000','24,000'),
('Optional Jev provision','1,000','1,000','12,000'),
('Static frontend DNS CDN and TLS','0','0','0'),
('Subtotal before tax','36,515','53,667',money(tot['ai']+tot['other'])),
]
table(['Cost category','2026 month','2027 month','Year total'],costs,[3.3,1.15,1.15,1.25],9.2)
p('Hosting: one API service and one worker at USD 25 each per month. Redis is provisioned at USD 10. These are initial capacity allowances based on Render service examples; they are not a throughput guarantee or a high-availability deployment. [2]')
p('Database: USD 25 Pro + USD 15 Small compute − USD 10 compute credit = USD 30 monthly. The base stays within 8 GB database disk and 250 GB uncached egress per month. Corpus growth, chat retention or retrieval traffic can exceed those quotas. [3]')
p('Storage: assume 50 GB of originals and 50 GB of backup copies. At R2 standard rates, storage itself is approximately INR 122 monthly after a 10 GB free allowance. The rest of the INR 1,000 envelope covers operations, traffic elsewhere and backup tooling. It does not buy a managed point-in-time recovery service. [5]')
p('Subscriptions: INR 1,000 monitoring, INR 2,000 communications and INR 2,000 business tools are provisional caps, not confirmed vendor subscriptions. Communications allow roughly 1,000 SMS, 4,000 WhatsApp utility messages, 6,000 emails and the WhatsApp platform fee. Marketing messages are excluded. [6]')
p('Optional Jev: the repo defaults its modes to off. The provision covers about INR 817 at 27,000 questions if aggregate decision input averages 8,000 tokens per question at USD 0.042 per million. Its inclusion is conservative and assumes no reduction in Gemini cost. [7, 10]')

# Page 5
newpage();h('4 Monthly operating cash projection')
p('The projection covers 1 November 2026 through 31 October 2027. Other costs include all non-generation categories on page 4, totalling INR 15,725 monthly. Cash need includes the provisional tax allowance and excludes contingency and initial onboarding.')
table(['Month','Days','Questions','AI cost','Other costs','Tax','Cash need'],[
    [r['month'],r['days'],money(r['queries']),money(r['ai']),money(r['other']),money(r['tax']),money(r['cash'])] for r in rows
]+[['Year',365,money(tot['queries']),money(tot['ai']),money(tot['other']),money(tot['tax']),money(tot['cash'])]],[1.0,.42,.9,1.1,1.12,1.05,1.16],9.0)
h2('How to read the monthly movement')
p('November to December changes reflect calendar length. The step in January primarily reflects the announced Flash price increase. February is lower because it contains 28 days. Later changes reflect calendar length; no price inflation or user growth has been assumed.')
p('Generation cost = question volume × applicable per-question rate. Other costs = INR 15,725. Tax allowance = 18% × (generation + other costs). Monthly cash need = generation + other costs + tax allowance.')
h2('Recurring spend composition')
table(['Annual measure','INR'],[
    ['AI generation before tax',money(tot['ai'])],
    ['All other recurring services and allowances before tax',money(tot['other'])],
    ['Provisional tax allowance',money(tot['tax'])],
    ['Average monthly cash requirement',money(tot['cash']/12)],
    ['Cash requirement per registered user per month',money(tot['cash']/12/1000)],
],[4.9,1.95],10)
p('Amounts are calculated at full precision and displayed to the nearest rupee. Displayed rows can differ from their displayed total by a few rupees because totals are rounded only after summation. Service subscriptions and annual renewals may also have different invoice dates from these monthly allocations.',size=10)

# Page 6
newpage();h('5 Funding schedule and initial costs')
p('Contingency is 15% of recurring cash need. The recommended allocation rounds the monthly funding requirement to INR 50,000 through December and INR 75,000 thereafter. An additional INR 10,000 envelope is assigned to November.')
cum=D(0);data=[]
for i,r in enumerate(rows):
    allocation=r['allocation']+(10000 if i==0 else 0);cum+=allocation
    data.append([r['month'],money(r['cash']),money(r['reserve']),money(allocation),money(cum)])
table(['Month','Cash need','Contingency','Allocation','Cumulative allocation'],data+[['Year',money(tot['cash']),money(tot['reserve']),'860,000','860,000']],[1.15,1.35,1.35,1.35,1.55],9.3)
h2('Initial provisioning envelope')
table(['One time item','INR cap'],[
    ['OCR and initial document embeddings including tax allowance','1,000'],
    ['Conditional sender registration or communication onboarding','6,000'],
    ['Other initial vendor fees and provisioning headroom','3,000'],
    ['Total initial envelope','10,000'],
],[5.0,1.85],10)
p('At the assumed 10,000-page corpus and 50% OCR share, initial OCR and embeddings are approximately INR 805 before tax if OCR free allowances are ignored. The INR 1,000 cap covers that estimate. Sender registration and other provisioning amounts are conditional budget provisions, not quoted fees or authorised purchases. Domain cost is already included in recurring allocations and is not charged again here.')
h2('Calendar adjustment and funding margin')
p(f'A simplified plan using twelve 30-day months would give INR 719,443 recurring cash need and INR 827,360 after 15% contingency. Actual calendar days increase those values to INR {money(tot["cash"])} and INR {money(tot["protected"])} respectively. The recommended allocation of INR 850,000 for recurring needs plus INR 10,000 for initial costs provides a further INR {money(D(850000)-tot["protected"])} above the calculated requirement.')

# Page 7
newpage();h('6 Scenarios and spending sensitivities')
p('Usage scenarios change the average daily active population while holding the registered base at 1,000, question mix constant, and non-generation envelopes unchanged. They are directional estimates; higher usage may require additional infrastructure capacity.')
sc=[]
for label,dau in [('Low',100),('Base',200),('High',350)]:
    cash,fund=scenario(dau)
    sc.append([label,str(dau),money(D(dau)*D('4.5')*365),money(cash),money(fund)])
table(['Scenario','DAU','Year questions','Recurring cash','Funding with reserve and setup'],sc,[.75,.5,1.4,1.65,2.45],9.4)
p('The base case supports the INR 8.60 lakh allocation. The high case requires approximately '+lakh(scenario(350)[1])+' including contingency and the initial envelope. The scenarios do not predict how many registered users will become subscribers.')
h2('Sensitivity around the base case')
sens=[]
for label,cash in [('4 questions per DAU daily',scenario(200,D(4))[0]),('5 questions per DAU daily',scenario(200,D(5))[0]),('25 percent higher generation cost',scenario(200,mult=D('1.25'))[0])]:
    sens.append([label,money(cash),money(cash-tot['cash'])])
table(['Change','Annual recurring cash','Change from base'],sens,[3.5,1.7,1.65],10)
p('Every additional 1,000 questions costs approximately INR 1,405 in generation from January, or INR 1,658 including the provisional tax allowance. This excludes any capacity step-up. A 10% currency increase would add at most about INR '+money(tot['cash']*D('.10'))+' annually if every cost moved with USD; actual exposure should be lower because some allowances are domestic.')
h2('Conditional costs outside the core total')
p('Payment processing depends on gross online collections. Use 3% of receipts including customer GST as a provisional collection-cost envelope, then replace it with the chosen gateway and recurring-mandate quote. For illustration, 150 subscriptions at INR 999 plus an assumed 18% customer GST produce INR 176,823 receipts and approximately INR 5,305 monthly collection costs. No revenue or gateway fee is embedded in the base cost plan. [8]')
p('Managed point-in-time database recovery starts at USD 100 monthly, approximately INR 9,000 before tax, in addition to this base plan. A database upgrade from Small to Medium adds USD 45 monthly, or INR 4,050 before tax, assuming the same plan and credit. [3]')
p('Paid legal-content licences, advanced document extraction, web-search grounding, voice, satellite imagery, enterprise SSO, contractual availability guarantees, dedicated or restricted hosting, security audits and procurement compliance services require a defined scope and supplier quote before being added. None is assumed for the existing text-document Q&A workload.')

# Page 8
newpage();h('7 Budget controls and source register')
h2('Validation and review cadence')
p('During the first month, compare metered provider usage with question volume and record model costs by simple, procedural and complex routes. Review active users, billable reasoning tokens, retries, OCR pages, database size, bandwidth and message counts. Use those measurements to replace the largest assumptions before the second monthly budget review.')
p('Review the forecast if monthly question volume exceeds 30,000, average generation cost exceeds the applicable allowance by 20%, database disk reaches 80% of included capacity, or a required government procurement term changes the hosting or assurance scope. Set provider billing alerts and apply per-user usage allowances before offering broad commercial access.')
h2('Source register')
p('Public pricing checked on 15 October 2026. Links identify the published rate or product reference used; estimates labelled as allowances remain internal planning assumptions.',size=10)
sources=[
('[1] Google Gemini API pricing including text generation and embeddings','https://ai.google.dev/gemini-api/docs/pricing'),
('[2] Render maintained deployment example with service and queue costs','https://github.com/render-examples/dify-render-template'),
('[3] Supabase pricing compute allowances backups and optional recovery','https://supabase.com/pricing'),
('[4] Google Cloud Document AI OCR pricing','https://cloud.google.com/products/document-ai/pricing'),
('[5] Cloudflare R2 storage and operation pricing','https://developers.cloudflare.com/r2/pricing/'),
('[6] MSG91 WhatsApp platform subscription and message pricing reference','https://msg91.com/help/whatsapp/whatsapp-subscription/'),
('[7] TypeSafe AI published Jev input token rate','https://typesafe.ai/blog/introducing-system-one-models-and-jev?render=full'),
('[8] Razorpay payment gateway pricing','https://razorpay.com/pricing/'),
('[9] Cloudflare free plan services','https://www.cloudflare.com/plans/'),
]
for label,url in sources:link(label,url)
p('[10] Project evidence: README.md describes the retrieval architecture; settings.py and services/gemini.py identify model defaults; generation_cost.py supplies the existing generation cost conventions; docs/cost_policy.md records selected evaluation runs; jev_settings.py and docs/jev.md describe optional Jev modes. These are repository references rather than live production measurements.',size=9.5)
h2('Decision basis')
p('Approve an INR 860,000 first-year operating envelope for 1,000 registered users, subject to the stated activity, question-mix and standard-cloud assumptions. Treat the INR 10,000 initial allowance and optional integrations as controlled provisions. Development and maintenance labour remain outside this approval scope.')

path=OUT/'Aranyabodh Annual Operating Cost Projection.docx'
doc.save(path)
serial=lambda v:float(v) if isinstance(v,Decimal) else v
(OUT/'projection_calculations.json').write_text(json.dumps({'months':rows,'totals':tot,'scenarios':sc},default=serial,indent=2))
assert sum(r['days'] for r in rows)==365
assert tot['queries']==328500
assert tot['allocation']+10000==860000
assert sum(D(x) for x in ['2250','2250','2700','900','1000','1000','500','2000','125','2000','1000'])==FIXED
print(path)
print(json.dumps(tot,default=serial))
