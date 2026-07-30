import re
import struct
from collections import Counter
from collections.abc import Iterator
from pathlib import Path

import olefile
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pypdf import PdfReader

from rag_errors import RagError
from services.document_storage import SUPPORTED_STORAGE_EXTENSIONS, document_storage
from settings import DOCS_DIR as _DOCS_DIR
from settings import env_bool, env_int

SUPPORTED_EXTENSIONS = SUPPORTED_STORAGE_EXTENSIONS
DOCS_DIR = _DOCS_DIR
GENERIC_TITLE_LINES = {
    "EXTRAORDINARY",
    "GAZETTE OF INDIA",
    "THE GAZETTE OF INDIA",
    "PUBLISHED BY AUTHORITY",
    "NOTIFICATION",
    "RESOLUTION",
}


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def read_txt(path: Path) -> list[dict]:
    text = normalize_text(path.read_text(encoding="utf-8"))
    return [{"page": None, "text": text}] if text else []


def read_pdf(path: Path) -> list[dict]:
    if not env_bool("PDF_EXTRACT_TABLES", True):
        return remove_repeated_margin_lines(read_pdf_with_pypdf(path))

    try:
        pages = read_pdf_with_pdfplumber(path)
        if pages:
            return remove_repeated_margin_lines(pages)
    except Exception:
        pass

    return remove_repeated_margin_lines(read_pdf_with_pypdf(path))


def read_pdf_with_pypdf(path: Path) -> list[dict]:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise RagError(f"Could not read PDF {path.name}: {exc}") from exc

    enforce_pdf_page_limit(path, len(reader.pages))
    pages = []
    extracted_chars = 0
    for i, page in enumerate(reader.pages, start=1):
        text = normalize_text(page.extract_text() or "")
        if text:
            extracted_chars += len(text)
            enforce_extracted_text_limit(path, extracted_chars)
            pages.append({"page": i, "text": text})
    return pages


def read_pdf_with_pdfplumber(path: Path, *, extract_tables: bool = True) -> list[dict]:
    import pdfplumber

    pages = []
    table_index = 0
    extracted_chars = 0
    with pdfplumber.open(str(path)) as pdf:
        enforce_pdf_page_limit(path, len(pdf.pages))
        for page_number, page in enumerate(pdf.pages, start=1):
            try:
                blocks = []
                text = normalize_text(page.extract_text() or "")
                if text:
                    blocks.append({"type": "text", "text": text})

                if extract_tables:
                    for table in page.extract_tables():
                        table_block = table_block_from_rows(table, table_index)
                        table_index += 1
                        if table_block:
                            blocks.append(table_block)

                page_text = normalize_text("\n\n".join(block["text"] for block in blocks))
                if page_text:
                    extracted_chars += len(page_text)
                    enforce_extracted_text_limit(path, extracted_chars)
                    pages.append({"page": page_number, "text": page_text, "blocks": blocks})
            finally:
                close_page = getattr(page, "close", None)
                if callable(close_page):
                    close_page()
    return pages


def enforce_pdf_page_limit(path: Path, page_count: int) -> None:
    max_pages = env_int("MAX_PDF_PAGES", 500)
    if page_count > max_pages:
        raise RagError(f"PDF {path.name} has {page_count} pages; the configured limit is {max_pages}.")


def enforce_extracted_text_limit(path: Path, extracted_chars: int) -> None:
    max_chars = env_int("MAX_EXTRACTED_CHARS", 15_000_000)
    if extracted_chars > max_chars:
        raise RagError(f"Document {path.name} exceeds the extracted text limit of {max_chars} characters.")


def iter_docx_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def clean_cell_text(text: str) -> str:
    return normalize_text(text).replace("\n", " ")


def table_block_from_rows(raw_rows: list[list[str | None]], table_index: int) -> dict | None:
    rows = []
    for raw_row in raw_rows:
        cells = [clean_cell_text(cell or "") for cell in raw_row]
        if any(cells):
            rows.append(cells)

    if not rows:
        return None

    headers = rows[0] if len(rows) > 1 else []
    data_rows = rows[1:] if headers else rows
    rendered_rows = [" | ".join(cell for cell in row) for row in rows]
    return {
        "type": "table",
        "table_index": table_index,
        "headers": headers,
        "rows": data_rows,
        "text": normalize_text("\n".join(rendered_rows)),
    }


def docx_table_block(table: Table, table_index: int) -> dict | None:
    return table_block_from_rows([[cell.text for cell in row.cells] for row in table.rows], table_index)


def read_docx(path: Path) -> list[dict]:
    try:
        document = Document(str(path))
    except Exception as exc:
        raise RagError(f"Could not read DOCX {path.name}: {exc}") from exc

    blocks = []
    table_index = 0
    for block in iter_docx_blocks(document):
        if isinstance(block, Paragraph):
            text = normalize_text(block.text)
            if text:
                blocks.append({"type": "text", "text": text})
        elif isinstance(block, Table):
            table_block = docx_table_block(block, table_index)
            table_index += 1
            if table_block:
                blocks.append(table_block)

    text = normalize_text("\n\n".join(block["text"] for block in blocks))
    return [{"page": None, "text": text, "blocks": blocks}] if text else []


def iter_pptx_shapes(shapes):
    for shape in shapes:
        if getattr(shape, "shape_type", None) == 6:  # MSO_SHAPE_TYPE.GROUP
            yield from iter_pptx_shapes(shape.shapes)
        else:
            yield shape


def read_pptx(path: Path) -> list[dict]:
    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise RagError(f"Could not read PPTX {path.name}: {exc}") from exc

    pages = []
    table_index = 0
    extracted_chars = 0
    for slide_number, slide in enumerate(presentation.slides, start=1):
        blocks = []
        for shape in iter_pptx_shapes(slide.shapes):
            if getattr(shape, "has_text_frame", False):
                text = normalize_text(shape.text_frame.text)
                if text:
                    blocks.append({"type": "text", "text": text})
            if getattr(shape, "has_table", False):
                table_block = table_block_from_rows(
                    [[cell.text for cell in row.cells] for row in shape.table.rows],
                    table_index,
                )
                table_index += 1
                if table_block:
                    blocks.append(table_block)

        text = normalize_text("\n\n".join(block["text"] for block in blocks))
        if text:
            extracted_chars += len(text)
            enforce_extracted_text_limit(path, extracted_chars)
            pages.append({"page": slide_number, "text": text, "blocks": blocks})
    return pages


def iter_ppt_records(data: bytes):
    offset = 0
    while offset + 8 <= len(data):
        options, record_type, length = struct.unpack_from("<HHI", data, offset)
        offset += 8
        end = offset + length
        if end > len(data):
            break
        yield options & 0xF, options >> 4, record_type, data[offset:end]
        offset = end


def legacy_ppt_text(payload: bytes) -> list[str]:
    texts = []
    for record_version, _record_instance, record_type, record_payload in iter_ppt_records(payload):
        if record_type == 4000:  # TextCharsAtom (UTF-16LE)
            text = record_payload.decode("utf-16-le", errors="replace")
            text = normalize_text(text.replace("\x00", ""))
            if text:
                texts.append(text)
        elif record_type == 4008:  # TextBytesAtom (compressed single-byte text)
            text = normalize_text(record_payload.decode("cp1252", errors="replace").replace("\x00", ""))
            if text:
                texts.append(text)
        elif record_version == 0xF:
            texts.extend(legacy_ppt_text(record_payload))
    return texts


def legacy_ppt_slide_lists(payload: bytes) -> list[list[str]]:
    slide_texts = []
    for record_version, record_instance, record_type, record_payload in iter_ppt_records(payload):
        if record_type == 4080 and record_version == 0xF and record_instance == 0:  # Slides (not masters or notes)
            current = None
            for child_version, child_instance, child_type, child_payload in iter_ppt_records(record_payload):
                if child_type == 1011:  # SlidePersistAtom
                    if current:
                        slide_texts.append(current)
                    current = []
                elif current is not None:
                    if child_type in {4000, 4008}:
                        child_options = child_version | (child_instance << 4)
                        current.extend(legacy_ppt_text(struct.pack("<HHI", child_options, child_type, len(child_payload)) + child_payload))
                    elif child_version == 0xF:
                        current.extend(legacy_ppt_text(child_payload))
            if current:
                slide_texts.append(current)
        elif record_version == 0xF:
            slide_texts.extend(legacy_ppt_slide_lists(record_payload))
    return slide_texts


def legacy_ppt_slide_containers(payload: bytes) -> list[list[str]]:
    slide_texts = []
    for record_version, _record_instance, record_type, record_payload in iter_ppt_records(payload):
        if record_type == 1006 and record_version == 0xF:  # SlideContainer
            texts = legacy_ppt_text(record_payload)
            if texts:
                slide_texts.append(texts)
        elif record_version == 0xF:
            slide_texts.extend(legacy_ppt_slide_containers(record_payload))
    return slide_texts


def read_ppt(path: Path) -> list[dict]:
    try:
        with olefile.OleFileIO(str(path)) as presentation:
            if not presentation.exists("PowerPoint Document"):
                raise ValueError("PowerPoint Document stream is missing")
            payload = presentation.openstream("PowerPoint Document").read()
    except Exception as exc:
        raise RagError(f"Could not read PPT {path.name}: {exc}") from exc

    slide_texts = legacy_ppt_slide_lists(payload) or legacy_ppt_slide_containers(payload)
    if not slide_texts:
        all_text = legacy_ppt_text(payload)
        slide_texts = [all_text] if all_text else []

    pages = []
    extracted_chars = 0
    for slide_number, texts in enumerate(slide_texts, start=1):
        text = normalize_text("\n\n".join(texts))
        if text:
            extracted_chars += len(text)
            enforce_extracted_text_limit(path, extracted_chars)
            pages.append({"page": slide_number, "text": text, "blocks": [{"type": "text", "text": text}]})
    return pages


def iter_documents(source: str | None = None) -> Iterator[dict]:
    with document_storage().document_files(source=source) as files:
        for document_file in files:
            doc = read_document(document_file)
            if doc:
                yield doc


def load_documents(source: str | None = None) -> list[dict]:
    return list(iter_documents(source=source))


def read_document(document_file) -> dict | None:
    path = document_file.path
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        return None

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        pages = read_pdf(path)
    elif suffix == ".docx":
        pages = read_docx(path)
    elif suffix == ".pptx":
        pages = read_pptx(path)
    elif suffix == ".ppt":
        pages = read_ppt(path)
    else:
        pages = read_txt(path)

    if not pages:
        return None

    title = infer_title(document_file.name, pages)
    return {
        "source": document_file.name,
        "kind": path.suffix.lower().lstrip("."),
        "title": title,
        "page_count": len([p for p in pages if p["page"] is not None]) or None,
        "metadata": extract_document_metadata(document_file.name, title, pages),
        "pages": pages,
    }


def infer_title(source: str, pages: list[dict]) -> str:
    candidates = []
    source_title = re.sub(r"^[\d_\s-]+", "", Path(source).stem).replace("_", " ").replace("-", " ").strip()
    source_has_title_keyword = re.search(
        r"\b(act|rules?|guidelines?|handbook|procedure|process|workflow|circular|order|notification|faq|checklist)\b",
        source_title,
        re.I,
    )
    if len(re.findall(r"[A-Za-z]", source_title)) >= 8 or (
        source_has_title_keyword and len(re.findall(r"[A-Za-z]", source_title)) >= 5
    ):
        source_score = 35
        if source_has_title_keyword:
            source_score += 40
        if source_title.islower():
            source_score -= 15
        candidates.append((source_score, source_title))

    for page in pages[:3]:
        for position, raw_line in enumerate(page["text"].splitlines()[:30]):
            line = normalize_text(raw_line)
            if not 8 <= len(line) <= 180 or is_generic_title_line(line):
                continue
            score = max(0, 20 - position)
            if re.search(
                r"\b(act|rules?|guidelines?|handbook|procedure|process|workflow|circular|order|notification|scheme|faq|check\s*list)\b",
                line,
                re.I,
            ):
                score += 40
            if re.search(r"\b(ministry|department|government of)\b", line, re.I):
                score += 30
            if re.search(r"\b(19|20)\d{2}\b", line):
                score += 5
            if line.isupper():
                score += 5
            if line[:1].islower():
                score -= 30
            if len(line) > 120:
                score -= 50
            candidates.append((score, line))
        if len(candidates) > 1:
            break

    if candidates:
        return max(candidates, key=lambda item: item[0])[1]
    return Path(source).stem.replace("_", " ").replace("-", " ").title()


def is_generic_title_line(line: str) -> bool:
    normalized = re.sub(r"\s+", " ", line).strip(" .-").upper()
    if normalized in GENERIC_TITLE_LINES:
        return True
    if re.match(r"^\d+\s+GI/\d{4}", normalized):
        return True
    if re.search(r"\b(REGD\.?\s*NO|REGISTERED\s*NO)\b", normalized):
        return True
    if re.match(r"^X{3}GID", normalized):
        return True
    if re.match(r"^(FILE|NO\.)\s*(NO\.?)?[:.\s-]*[A-Z0-9/_-]+$", normalized):
        return True
    return len(re.sub(r"[^A-Za-z]", "", normalized)) < 4


def extract_document_metadata(source: str, title: str, pages: list[dict]) -> dict:
    sample = normalize_text("\n".join(page["text"] for page in pages[:5]))[:20000]
    haystack = f"{title}\n{source}\n{sample}"
    document_type = infer_document_type(f"{title}\n{source}")
    if document_type == "document":
        document_type = infer_document_type(sample)
    return {
        "document_type": document_type,
        "identifiers": extract_legal_identifiers(haystack),
        "years": sorted(set(re.findall(r"\b(?:19|20)\d{2}\b", haystack))),
        "authority": infer_authority(sample),
    }


def infer_document_type(text: str) -> str:
    lowered = text.lower()
    types = (
        ("faq", r"\b(faq|frequently asked questions?)\b"),
        ("notification", r"\bnotification\b"),
        ("circular", r"\bcircular\b"),
        ("order", r"\border\b"),
        ("procedure", r"\b(procedure|process|workflow|sop)\b"),
        ("guidelines", r"\b(guidelines?|handbook|check\s*list)\b"),
        ("rules", r"\brules?,?\s*(?:19|20)\d{2}\b"),
        ("act", r"\bact,?\s*(?:19|20)\d{2}\b"),
    )
    for document_type, pattern in types:
        if re.search(pattern, lowered):
            return document_type
    return "document"


def infer_authority(text: str) -> str | None:
    for raw_line in text.splitlines()[:80]:
        line = normalize_text(raw_line)
        if 8 <= len(line) <= 180 and re.search(r"\b(ministry|department|government of|govt\.? of)\b", line, re.I):
            return line
    return None


def extract_legal_identifiers(text: str) -> list[str]:
    patterns = [
        r"\b(?:section|sec\.?|rule|rules?|article|schedule|annexure|appendix)\s*[-:]?\s*\d+(?:\.\d+)*(?:\s*\([a-zA-Z0-9ivx]+\))*",
        r"\bG\.?\s*S\.?\s*R\.?\s*[-.:]?\s*\d+\s*\([A-Za-z]\)",
        r"\bS\.?\s*R\.?\s*O\.?\s*[-.:]?\s*\d+",
        r"\b[A-Z]{1,6}-\d+(?:/\d+){1,4}/(?:19|20)\d{2}(?:-[A-Z]{1,6})?\b",
        r"\b\d+(?:\.\d+)?\s*(?:ha|hectares?)\b",
    ]
    identifiers = []
    seen = set()
    for pattern in patterns:
        for match in re.findall(pattern, text, re.I):
            normalized = re.sub(r"\s+", " ", match).strip(" .")
            key = re.sub(r"[^a-z0-9]+", "", normalized.lower())
            if key and key not in seen:
                seen.add(key)
                identifiers.append(normalized)
    return identifiers[:100]


def remove_repeated_margin_lines(pages: list[dict]) -> list[dict]:
    if len(pages) < 3:
        return pages

    margin_counts = Counter()
    for page in pages:
        lines = [line.strip() for line in page["text"].splitlines() if line.strip()]
        for line in set(lines[:3] + lines[-3:]):
            if len(line) <= 160:
                margin_counts[line] += 1

    threshold = max(3, int(len(pages) * 0.5))
    repeated = {line for line, count in margin_counts.items() if count >= threshold}
    if not repeated:
        return pages

    cleaned_pages = []
    for page in pages:
        cleaned = dict(page)
        cleaned["text"] = normalize_text("\n".join(line for line in page["text"].splitlines() if line.strip() not in repeated))
        cleaned["blocks"] = [
            {
                **block,
                "text": normalize_text(
                    "\n".join(line for line in block.get("text", "").splitlines() if line.strip() not in repeated)
                ),
            }
            if block.get("type") == "text"
            else block
            for block in page.get("blocks", [])
        ]
        if cleaned["text"]:
            cleaned_pages.append(cleaned)
    return cleaned_pages
